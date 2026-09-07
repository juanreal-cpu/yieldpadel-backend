import json
import logging
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional
import unicodedata

from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.competitor import CompetitorClub

logger = logging.getLogger("yieldpadel.radar_loader")

BASE_DIR = Path(__file__).resolve().parent.parent.parent
DATA_DIR = BASE_DIR / "app" / "data"
EXCEL_CANDIDATE_PATHS = [
    BASE_DIR / "data" / "directorio-clubes-padel-colombia.xlsx",
    BASE_DIR / "directorio-clubes-padel-colombia.xlsx",
    Path("data/directorio-clubes-padel-colombia.xlsx").resolve(),
]
JSON_CACHE_PATH = DATA_DIR / "radar_clubs.json"


def normalize_text(text: Optional[str]) -> str:
    """Normalize text removing accents and extra whitespace for robust matching."""
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    return text.lower().strip()


def assign_market_prices(name: str, city: str, zone: str, is_target: bool) -> tuple[float, float]:
    """
    Assign realistic market benchmark prices (Valle, Pico) in COP.
    Capital Pádel Maloka is target partner ($80.000 Valle / $120.000 Pico).
    """
    if is_target:
        return 80000.0, 120000.0

    c_norm = normalize_text(city)
    z_norm = normalize_text(zone)

    if "bogota" in c_norm:
        if "chico" in z_norm or "chapinero" in z_norm:
            return 90000.0, 140000.0
        elif any(k in z_norm for k in ["suba", "usaquen", "colina", "autonorte", "cedritos"]):
            return 85000.0, 130000.0
        else:
            return 85000.0, 130000.0

    if "medellin" in c_norm:
        if "poblado" in z_norm:
            return 90000.0, 140000.0
        return 85000.0, 130000.0

    if "cali" in c_norm:
        return 80000.0, 120000.0

    if "barranquilla" in c_norm:
        return 85000.0, 130000.0

    if "bucaramanga" in c_norm:
        return 75000.0, 110000.0

    if "cartagena" in c_norm:
        return 95000.0, 150000.0

    if "pereira" in c_norm:
        return 75000.0, 110000.0

    return 85000.0, 130000.0


def extract_clubs_from_excel() -> List[Dict[str, Any]]:
    """Parse the clubs directory from the Excel workbook."""
    excel_path = None
    for path in EXCEL_CANDIDATE_PATHS:
        if path.is_file():
            excel_path = path
            break

    if not excel_path:
        logger.warning("directorio-clubes-padel-colombia.xlsx not found on disk.")
        return []

    try:
        import openpyxl

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet_name = "Directorio Nacional" if "Directorio Nacional" in wb.sheetnames else wb.sheetnames[0]
        sheet = wb[sheet_name]

        clubs = []
        for r in range(9, sheet.max_row + 1):
            name_val = sheet.cell(r, 2).value
            if not name_val or "total" in str(name_val).lower():
                continue

            name = str(name_val).strip()
            city = str(sheet.cell(r, 3).value or "").strip()
            address = str(sheet.cell(r, 4).value or "").strip()
            zone = str(sheet.cell(r, 5).value or "").strip()
            courts_val = sheet.cell(r, 6).value
            rating_val = sheet.cell(r, 7).value
            phone = str(sheet.cell(r, 8).value or "").strip()
            website = str(sheet.cell(r, 9).value or "").strip()
            coords_val = sheet.cell(r, 10).value

            lat, lng = None, None
            if coords_val and "," in str(coords_val):
                parts = str(coords_val).split(",")
                try:
                    lat = float(parts[0].strip())
                    lng = float(parts[1].strip())
                except (ValueError, TypeError):
                    lat, lng = None, None

            try:
                courts_count = int(courts_val) if courts_val is not None else 4
            except (ValueError, TypeError):
                courts_count = 4

            try:
                rating = float(rating_val) if rating_val is not None else 4.5
            except (ValueError, TypeError):
                rating = 4.5

            # Target partner identifier: Capital Pádel Maloka (Salitre)
            is_target = ("maloka" in name.lower()) or ("capital padel maloka" in normalize_text(name))

            p_valle, p_pico = assign_market_prices(name, city, zone, is_target)

            clubs.append(
                {
                    "name": name,
                    "city": city,
                    "zone": zone,
                    "address": address,
                    "latitude": lat,
                    "longitude": lng,
                    "courts_count": courts_count,
                    "rating": rating,
                    "phone": phone,
                    "website": website,
                    "price_valle": p_valle,
                    "price_pico": p_pico,
                    "is_target_partner": is_target,
                }
            )

        # Save to JSON cache
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        with open(JSON_CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump(clubs, f, ensure_ascii=False, indent=2)

        return clubs
    except Exception as e:
        logger.error(f"Error parsing Excel directory: {e}")
        return []


def load_clubs_dataset() -> List[Dict[str, Any]]:
    """Load clubs dataset from Excel or fallback JSON cache."""
    clubs = extract_clubs_from_excel()
    if clubs:
        return clubs

    if JSON_CACHE_PATH.is_file():
        try:
            with open(JSON_CACHE_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.error(f"Error loading JSON cache {JSON_CACHE_PATH}: {e}")

    # Ultimate fallback hardcoded list for Capital Padel Maloka and key competitors
    return [
        {
            "name": "Capital Padel Maloka",
            "city": "Bogotá D.C.",
            "zone": "Salitre",
            "address": "Cra 69 #24a-67",
            "latitude": 4.656293,
            "longitude": -74.108702,
            "courts_count": 4,
            "rating": 4.8,
            "phone": "+57 302 4683239",
            "website": "https://instagram.com/capitalpadelclubbogota",
            "price_valle": 80000.0,
            "price_pico": 120000.0,
            "is_target_partner": True,
        },
        {
            "name": "Bogotá Pádel Club",
            "city": "Bogotá D.C.",
            "zone": "Chicó / Chapinero",
            "address": "Ak 7 #69-41",
            "latitude": 4.652345,
            "longitude": -74.057207,
            "courts_count": 5,
            "rating": 4.9,
            "phone": "+57 316 5306717",
            "website": "https://bogotapadelclub.com/",
            "price_valle": 90000.0,
            "price_pico": 140000.0,
            "is_target_partner": False,
        },
        {
            "name": "Combo Padel Club",
            "city": "Bogotá D.C.",
            "zone": "Suba",
            "address": "Cra. 76 #175-20",
            "latitude": 4.765632,
            "longitude": -74.064072,
            "courts_count": 4,
            "rating": 4.7,
            "phone": "",
            "website": "",
            "price_valle": 85000.0,
            "price_pico": 130000.0,
            "is_target_partner": False,
        },
    ]


async def ensure_competitor_clubs(db: AsyncSession) -> List[CompetitorClub]:
    """Ensure database has competitor clubs populated. Seed if empty."""
    res = await db.execute(select(func.count(CompetitorClub.id)))
    count = res.scalar() or 0

    if count == 0:
        dataset = load_clubs_dataset()
        for item in dataset:
            club = CompetitorClub(
                name=item["name"],
                city=item["city"],
                zone=item["zone"],
                address=item["address"],
                latitude=item["latitude"],
                longitude=item["longitude"],
                courts_count=item["courts_count"],
                rating=item["rating"],
                phone=item["phone"],
                website=item["website"],
                price_valle=item["price_valle"],
                price_pico=item["price_pico"],
                is_target_partner=item["is_target_partner"],
            )
            db.add(club)
        await db.commit()

    stmt = select(CompetitorClub).order_by(CompetitorClub.is_target_partner.desc(), CompetitorClub.name.asc())
    clubs_res = await db.execute(stmt)
    return list(clubs_res.scalars().all())
